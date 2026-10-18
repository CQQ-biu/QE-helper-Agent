"""根据问题关键词确定权威来源文件。"""
import csv
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import load_workbook


@dataclass(frozen=True)
class SourceRule:
    keyword: str
    source: str
    priority: int = 0


class SourceMappingResolver:
    """读取业务维护的关键词清单，并在文件变化时自动刷新。"""

    _keyword_fields = {"keyword", "keywords", "关键词", "关键字"}
    _source_fields = {"source", "source_file", "filename", "file", "来源", "来源文件", "文件名"}
    _priority_fields = {"priority", "优先级"}

    def __init__(self, mapping_path: Path):
        self.mapping_path = mapping_path
        self._fingerprint = None
        self._rules: List[SourceRule] = []

    def resolve(self, question: str) -> Optional[SourceRule]:
        self._reload_if_changed()
        normalized = question.casefold()
        matches = [rule for rule in self._rules if rule.keyword.casefold() in normalized]
        if not matches:
            return None
        # 优先级高的规则优先；同优先级下，长关键词比短关键词更具体。
        return sorted(matches, key=lambda rule: (-rule.priority, -len(rule.keyword), rule.source))[0]

    def _reload_if_changed(self) -> None:
        fingerprint = None
        if self.mapping_path.exists():
            stat = self.mapping_path.stat()
            fingerprint = (stat.st_mtime_ns, stat.st_size)
        if fingerprint == self._fingerprint:
            return
        self._rules = self._load_rules() if fingerprint else []
        self._fingerprint = fingerprint

    def _load_rules(self) -> List[SourceRule]:
        suffix = self.mapping_path.suffix.lower()
        if suffix == ".xlsx":
            records = self._read_xlsx()
        elif suffix == ".csv":
            with self.mapping_path.open(encoding="utf-8-sig", newline="") as handle:
                records = list(csv.DictReader(handle))
        elif suffix == ".json":
            data = json.loads(self.mapping_path.read_text(encoding="utf-8"))
            records = (
                [{"keyword": keyword, "source": source} for keyword, source in data.items()]
                if isinstance(data, dict) else data
            )
        else:
            raise ValueError("来源映射清单仅支持 .xlsx、.csv 或 .json")

        rules: List[SourceRule] = []
        for record in records:
            normalized = {str(key).strip().casefold(): value for key, value in record.items() if key is not None}
            keyword_value = self._field(normalized, self._keyword_fields)
            source_value = self._field(normalized, self._source_fields)
            priority_value = self._field(normalized, self._priority_fields) or 0
            if keyword_value is None or source_value is None:
                continue
            try:
                priority = int(priority_value)
            except (TypeError, ValueError):
                priority = 0
            for keyword in re.split(r"[|,，;；]", str(keyword_value)):
                keyword = keyword.strip()
                if keyword:
                    rules.append(SourceRule(keyword, str(source_value).strip(), priority))
        return rules

    def _read_xlsx(self) -> List[Dict]:
        workbook = load_workbook(self.mapping_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(rows, [])]
            return [dict(zip(headers, row)) for row in rows]
        finally:
            workbook.close()

    @staticmethod
    def _field(record: Dict, aliases: set):
        for alias in aliases:
            if alias.casefold() in record:
                return record[alias.casefold()]
        return None
