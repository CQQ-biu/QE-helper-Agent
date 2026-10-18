"""RAG实现：知识库解析、递归分块、pgvector检索和Agent受约束生成。"""
import csv
import json
import re
import urllib.error
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from openpyxl import load_workbook

from app.services.recursive_chunker import RecursiveSemanticChunker
from app.services.source_mapping import SourceMappingResolver, SourceRule
from app.services.vector_store import PostgresVectorStore
from app.utils.logger import logger


@dataclass(frozen=True)
class Chunk:
    text: str
    source: str


@dataclass(frozen=True)
class RAGResult:
    answer: str
    similarity: float
    sources: List[str]
    iterations: int
    search_queries: List[str]
    source_rule_keyword: Optional[str] = None


class OpenAICompatibleClient:
    """仅依赖标准库的 OpenAI 兼容 HTTP 客户端。"""

    def __init__(self, base_url: str, api_key: str, timeout: int):
        self.base_url = base_url.rstrip("/")
        self.api_key = api_key
        self.timeout = timeout

    def _post(self, path: str, payload: dict) -> dict:
        request = urllib.request.Request(
            f"{self.base_url}{path}",
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            headers={"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                return json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            raise RuntimeError(f"RAG接口请求失败({exc.code}): {detail[:500]}") from exc

    def embed(self, texts: Sequence[str], model: str) -> List[List[float]]:
        data = self._post("/embeddings", {"model": model, "input": list(texts)})
        return [item["embedding"] for item in sorted(data["data"], key=lambda item: item["index"])]

    def complete(
        self,
        messages: list,
        model: str,
        temperature: float,
        tools: Optional[list] = None,
        tool_choice: Optional[str] = None,
    ) -> dict:
        payload: Dict[str, Any] = {
            "model": model,
            "messages": messages,
            "temperature": temperature,
        }
        if tools:
            payload["tools"] = tools
        if tool_choice is not None:
            payload["tool_choice"] = tool_choice
        data = self._post("/chat/completions", payload)
        return data["choices"][0]["message"]


class RAGService:
    """使用PostgreSQL/pgvector持久化索引的RAG Agent服务。"""

    def __init__(self, settings, client=None, vector_store=None):
        backend_dir = Path(__file__).resolve().parents[2]
        kb_path = Path(settings.RAG_KNOWLEDGE_BASE_DIR)
        self.knowledge_dir = kb_path if kb_path.is_absolute() else (backend_dir / kb_path).resolve()
        mapping_setting = getattr(settings, "RAG_SOURCE_MAPPING_FILE", "")
        mapping_path = Path(mapping_setting) if mapping_setting else None
        if mapping_path and not mapping_path.is_absolute():
            mapping_path = (backend_dir / mapping_path).resolve()
        self.source_resolver = SourceMappingResolver(mapping_path) if mapping_path else None
        self.embedding_model = settings.RAG_EMBEDDING_MODEL
        self.chat_model = settings.RAG_CHAT_MODEL
        self.top_k = settings.RAG_TOP_K
        self.threshold = settings.RAG_SIMILARITY_THRESHOLD
        self.chunk_size = settings.RAG_CHUNK_SIZE
        self.chunk_overlap = settings.RAG_CHUNK_OVERLAP
        self.batch_size = settings.RAG_EMBEDDING_BATCH_SIZE
        self.temperature = settings.RAG_TEMPERATURE
        self.max_iterations = settings.RAG_AGENT_MAX_ITERATIONS
        if self.chunk_overlap >= self.chunk_size:
            raise ValueError("RAG_CHUNK_OVERLAP 必须小于 RAG_CHUNK_SIZE")
        if self.chunk_size <= 0 or self.chunk_overlap < 0:
            raise ValueError("RAG分块参数必须为正数，且重叠长度不能为负数")
        if self.top_k <= 0 or self.batch_size <= 0:
            raise ValueError("RAG_TOP_K 和 RAG_EMBEDDING_BATCH_SIZE 必须大于0")
        if self.max_iterations <= 0:
            raise ValueError("RAG_AGENT_MAX_ITERATIONS 必须大于0")
        if not 0 <= self.threshold <= 1:
            raise ValueError("RAG_SIMILARITY_THRESHOLD 必须在0到1之间")
        self.chunker = RecursiveSemanticChunker(self.chunk_size, self.chunk_overlap)
        if client is None:
            if not settings.RAG_API_KEY:
                raise ValueError("未配置 RAG_API_KEY，请在 backend/.env 中设置")
            client = OpenAICompatibleClient(settings.RAG_BASE_URL, settings.RAG_API_KEY, settings.RAG_REQUEST_TIMEOUT)
        self.client = client
        self.vector_store = vector_store or PostgresVectorStore(
            settings, self.client, self.knowledge_dir
        )

    def answer(self, question: str, language: str) -> Optional[RAGResult]:
        self._synchronize_index()
        if not self.vector_store.has_chunks():
            logger.warning("知识库为空: %s", self.knowledge_dir)
            return None

        source_rule = self.source_resolver.resolve(question) if self.source_resolver else None
        target_language = "简体中文" if language == "简体中文" else "English"
        routing_instruction = (
            f"业务规则已将本问题路由到文件“{source_rule.source}”（命中关键词“{source_rule.keyword}”）。"
            "检索工具会自动限制在该文件内，不得改用其他来源。"
            if source_rule else "本问题未命中来源映射规则，按向量检索结果确定来源。"
        )
        messages = [
            {
                "role": "system",
                "content": (
                    "你是质量稽核问答 Agent。你必须先调用 search_knowledge 搜索知识库。"
                    "分析搜索结果是否足够；如果不足，可改写关键词再次搜索。"
                    "只能依据工具返回的资料回答，不得使用外部知识或猜测。"
                    "证据不足时只回答‘未检索到’；证据充分时给出准确、简洁、可直接填入稽核表的答案。"
                    f"{routing_instruction}"
                    f"使用{target_language}回答。"
                ),
            },
            {"role": "user", "content": question},
        ]
        tools = [
            {
                "type": "function",
                "function": {
                    "name": "search_knowledge",
                    "description": "按语义搜索质量稽核知识库。可使用不同关键词多次调用。",
                    "parameters": {
                        "type": "object",
                        "properties": {
                            "query": {"type": "string", "description": "适合知识库检索的完整查询"}
                        },
                        "required": ["query"],
                    },
                },
            }
        ]
        best_similarity = 0.0
        sources: List[str] = []
        search_queries: List[str] = []

        for iteration in range(1, self.max_iterations + 1):
            message = self.client.complete(
                messages,
                self.chat_model,
                self.temperature,
                tools=tools,
                tool_choice="required" if iteration == 1 else "auto",
            )
            tool_calls = message.get("tool_calls") or []
            if not tool_calls:
                answer = (message.get("content") or "").strip()
                return self._build_result(
                    answer, best_similarity, sources, iteration, search_queries, source_rule
                )

            messages.append({
                "role": "assistant",
                "content": message.get("content"),
                "tool_calls": tool_calls,
            })
            for tool_call in tool_calls:
                tool_result, similarity, found_sources, query = self._execute_tool(
                    tool_call, question, source_rule
                )
                best_similarity = max(best_similarity, similarity)
                sources.extend(source for source in found_sources if source not in sources)
                search_queries.append(query)
                messages.append({
                    "role": "tool",
                    "tool_call_id": tool_call.get("id", f"search-{iteration}"),
                    "name": "search_knowledge",
                    "content": tool_result,
                })

        # 达到循环上限后禁用工具，要求 Agent 基于已有观察结果强制收口。
        messages.append({
            "role": "system",
            "content": "已达到最大检索轮数。请立即依据已有工具结果给出最终答案；证据不足只回答‘未检索到’。",
        })
        final_message = self.client.complete(
            messages, self.chat_model, self.temperature, tools=tools, tool_choice="none"
        )
        answer = (final_message.get("content") or "").strip()
        return self._build_result(
            answer, best_similarity, sources, self.max_iterations, search_queries, source_rule
        )

    def _execute_tool(
        self,
        tool_call: dict,
        fallback_query: str,
        source_rule: Optional[SourceRule],
    ) -> Tuple[str, float, List[str], str]:
        function = tool_call.get("function") or {}
        if function.get("name") != "search_knowledge":
            return json.dumps({"error": "未知工具"}, ensure_ascii=False), 0.0, [], fallback_query
        try:
            arguments = json.loads(function.get("arguments") or "{}")
        except json.JSONDecodeError:
            arguments = {}
        query = str(arguments.get("query") or fallback_query).strip()
        ranked = self._search(query, source_rule)
        if not ranked or ranked[0][0] < self.threshold:
            routed_file_exists = self.vector_store.source_exists(source_rule.source) if source_rule else None
            payload = {
                "query": query,
                "found": False,
                "routed_source": source_rule.source if source_rule else None,
                "message": (
                    "来源映射指定的文件不在知识库索引中，请检查清单文件名。"
                    if source_rule and not routed_file_exists
                    else "没有达到相关度阈值的知识，请改写查询或最终回答未检索到。"
                ),
            }
            return json.dumps(payload, ensure_ascii=False), 0.0, [], query

        matches = [
            {
                "source": source_rule.source if source_rule else chunk.source,
                "indexed_source": chunk.source,
                "similarity": round(similarity, 4),
                "content": chunk.text,
            }
            for similarity, chunk in ranked if similarity >= self.threshold
        ]
        found_sources = (
            [source_rule.source] if source_rule
            else list(dict.fromkeys(match["source"] for match in matches))
        )
        payload = {
            "query": query,
            "found": True,
            "routed_by_keyword": source_rule.keyword if source_rule else None,
            "matches": matches,
        }
        return json.dumps(payload, ensure_ascii=False), ranked[0][0], found_sources, query

    def _search(
        self,
        query: str,
        source_rule: Optional[SourceRule] = None,
    ) -> List[Tuple[float, Chunk]]:
        query_vector = self.client.embed([query], self.embedding_model)[0]
        results = self.vector_store.search(
            query_vector,
            self.top_k,
            routed_source=source_rule.source if source_rule else None,
        )
        return [(result.similarity, Chunk(result.content, result.source)) for result in results]

    def _build_result(
        self,
        answer: str,
        similarity: float,
        sources: List[str],
        iterations: int,
        search_queries: List[str],
        source_rule: Optional[SourceRule] = None,
    ) -> Optional[RAGResult]:
        no_answer = answer.rstrip("。.").lower() in {"", "未检索到", "not found"}
        # 即使模型提前输出答案，也必须有一次达到阈值的真实检索作为证据。
        if no_answer or similarity < self.threshold or not sources:
            return None
        return RAGResult(
            answer, similarity, sources, iterations, search_queries,
            source_rule.keyword if source_rule else None,
        )

    def _synchronize_index(self) -> None:
        files = self._knowledge_files()
        self.vector_store.synchronize(files, self._chunk_texts)

    def _knowledge_files(self) -> List[Path]:
        if not self.knowledge_dir.exists():
            return []
        supported = {".txt", ".md", ".csv", ".json", ".xlsx", ".pdf", ".docx"}
        return sorted(path for path in self.knowledge_dir.rglob("*") if path.suffix.lower() in supported)

    def _chunk_texts(self, path: Path) -> List[str]:
        return self.chunker.split(self._read_file(path))

    @staticmethod
    def _read_file(path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="replace")
        if suffix == ".csv":
            with path.open(encoding="utf-8-sig", errors="replace", newline="") as handle:
                return "\n\n".join(" | ".join(row) for row in csv.reader(handle))
        if suffix == ".json":
            return json.dumps(json.loads(path.read_text(encoding="utf-8")), ensure_ascii=False, indent=2)
        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sections = []
                for sheet in workbook.worksheets:
                    rows = [
                        " | ".join(str(value) for value in row if value is not None)
                        for row in sheet.iter_rows(values_only=True)
                    ]
                    sections.append(f"# Sheet: {sheet.title}\n\n" + "\n\n".join(row for row in rows if row))
                return "\n\n".join(sections)
            finally:
                workbook.close()
        if suffix == ".pdf":
            from pypdf import PdfReader
            return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
        if suffix == ".docx":
            from docx import Document
            paragraphs = []
            for paragraph in Document(str(path)).paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = paragraph.style.name if paragraph.style else ""
                heading = re.match(r"Heading\s+(\d+)", style_name, re.IGNORECASE)
                if heading:
                    text = f"{'#' * min(int(heading.group(1)), 6)} {text}"
                paragraphs.append(text)
            return "\n\n".join(paragraphs)
        return ""

